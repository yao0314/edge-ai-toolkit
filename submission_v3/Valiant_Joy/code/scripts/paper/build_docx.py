"""Build paper/TAAI2026_submission.docx from paper/manuscript.md. Rewritten from scratch.

The previous version was clever and wrong in two ways that took three rounds to find, so this one
is deliberately dull:

  * It never sets EXACT line spacing. The old version gave the Normal style a fixed 12 pt lead,
    which is right for body text and catastrophic for pictures: Word clips an inline image to the
    line height, so every 60 mm figure rendered as a 12 pt sliver. Figures here sit in paragraphs
    with default spacing and nothing else.
  * It does not infer where a figure belongs from where the prose happens to mention it. That rule
    starved one figure completely and the leftover pass then dumped it after the reference list.
    Each figure is pinned to a named section; that is the whole algorithm.

Page geometry follows Springer LNCS A4 (122 x 193 mm text block), which is what TAAI 2026's
International Track requires. This is not the official Springer template - take that from the link
in the CFP before submitting - but the measure and the type size match, so the text will not
re-flow much when it moves across.

    python scripts/paper/build_docx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

SRC = ROOT / "paper" / "manuscript.md"
FIGS = ROOT / "paper" / "figs"
OUT = ROOT / "paper" / "TAAI2026_submission.docx"

PAGE_W, PAGE_H, TEXT_W, TEXT_H = 210, 297, 122, 193
FIG_W = 112                      # mm, narrower than the text block so it can never overflow

CAPTIONS = {
    1: ("fig1_version_sweep.png",
        "Speedup of INT8 over the FP32 sibling across eleven ONNX Runtime releases, measured on "
        "byte-identical quantized artifacts. The static path falls in two steps; the first costs "
        "24% at four threads and nothing at one. The dynamic path does not move."),
    2: ("fig2_method_by_class.png",
        "Which quantization method wins, by model class. Each row joins the two methods for one "
        "model, so the direction of the line is the answer. The axis is logarithmic because "
        "dynamic quantization applied to convolutions lands near 0.05x."),
    3: ("fig3_config_interaction.png",
        "Configuration interaction on MobileNetV2. Either factor alone is acceptable; their "
        "intersection is not. A one-factor-at-a-time sweep never visits the failing cell."),
    4: ("fig4_accuracy_by_family.png",
        "Accuracy change by model and method. The axis is clipped at -5.6 pt so that the small "
        "bars stay legible; the clipped bar carries its true value."),
}

# (section-number prefix, paragraphs into that section, kind, number)
PLAN = [
    ("IV.", 1, "fig", 1),
    ("V.", 1, "fig", 2),
    ("V.", 4, "tab", 1),
    ("V.", 8, "fig", 4),
    ("VI.", 1, "fig", 3),
    ("VII.", 1, "tab", 2),
]

MECHS = [
    ("Q/DQ boundary density",
     "A correlation between int8/FP32 crossings and latency, counted on the graph that executes.",
     "The counts came from the stored artifact, but the runtime fuses about 90% of those nodes "
     "before the first inference (416/599 to 40/97 on the hybrid model, 74/182 to 1/1 on "
     "ResNet-50). Recounted on the executed graph, the faster family runs more surviving crossings "
     "than the slower one."),
    ("Fusion blocking",
     "Restricting quantization to Conv and MatMul should be faster, because the elementwise "
     "operations become fusible again.",
     "It was 8-9% slower on the hybrid model. Per-node profiling does show the static path "
     "spending 27.5-28.5% of node time on decomposed elementwise chains its dynamic sibling never "
     "runs, so the phenomenon is real; it is not what governs the latency."),
    ("Memory-bound elementwise",
     "Restricting quantization should raise the fitted parallel fraction.",
     "It fell, 0.430 to 0.366, with serial time unchanged at 4.177 and 4.215 ms. The gain that "
     "does exist is at one thread and vanishes at four, which is the opposite of a bandwidth "
     "account."),
]


def shade(cell, hexv):
    el = cell._tc.get_or_add_tcPr()
    el.append(el.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): hexv}))


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(PAGE_W), Mm(PAGE_H)
    s.left_margin = s.right_margin = Mm((PAGE_W - TEXT_W) / 2)
    s.top_margin = s.bottom_margin = Mm((PAGE_H - TEXT_H) / 2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Deliberately no line_spacing: an exact point value clips every inline image to that height.


def para(doc, text, size=10, bold=False, italic=False, align=None, before=0, after=0, indent=0):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Mm(indent)
    return p


def figure(doc, n):
    png, caption = CAPTIONS[n]
    src = FIGS / png
    if not src.exists():
        print(f"  MISSING {src}", file=sys.stderr)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(src), width=Mm(FIG_W))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    c.paragraph_format.space_after = Pt(11)
    r = c.add_run(f"Fig. {n}. ")
    r.bold, r.font.size = True, Pt(9)
    r2 = c.add_run(caption)
    r2.font.size = Pt(9)


def grid(doc, headers, rows, widths, caption, number, small=9):
    para(doc, f"Table {number}. {caption}", size=9, before=11, after=4)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h, w in zip(t.rows[0].cells, headers, widths):
        c.width = Mm(w)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(small)
        shade(c, "E9EEF3")
    for row in rows:
        cells = t.add_row().cells
        for i, (c, v) in enumerate(zip(cells, row)):
            c.width = Mm(widths[i])
            c.text = ""
            run = c.paragraphs[0].add_run(str(v))
            run.font.size = Pt(small)
            if i >= 2 and any(ch.isdigit() for ch in str(v)) and len(str(v)) < 9:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "", after=8)


def table_one(doc):
    df = pd.read_csv(FIGS / "_fig_data.csv")
    order = ["ResNet-50", "MobileNetV2", "YOLOv8n", "1-D conv AE", "DeiT-tiny", "MobileViT"]
    df["k"] = df["model"].apply(lambda m: order.index(m) if m in order else 99)
    df = df.sort_values(["k", "method"], ascending=[True, False])
    rows = [[r["model"], "static QDQ" if r["method"] == "static" else "dynamic",
             f"{r['sp4']:.3f}", f"{r['drop']:+.2f}", f"{r['size'] * 100:.0f}%"]
            for _, r in df.iterrows()]
    grid(doc, ["Model", "Method", "Speedup (4T)", "Accuracy (pt)", "Size"], rows,
         [30, 24, 24, 24, 20],
         "Quantization method by model class: speed at four threads relative to the FP32 sibling, "
         "accuracy change in percentage points, and compressed size as a fraction of FP32. "
         "ONNX Runtime 1.30.0.", 1)


def table_two(doc):
    grid(doc, ["Mechanism", "Falsification condition, fixed in advance", "What was measured"],
         [[n, cond, out] for n, cond, out in MECHS], [26, 44, 52],
         "Three candidate mechanisms for the static-path slowdown. A condition that would falsify "
         "each was fixed before taking the measurement that could fire it. All three fired.",
         2, small=8)


def main() -> int:
    if not SRC.exists():
        print(f"{SRC} not found", file=sys.stderr)
        return 1
    lines = [ln.rstrip() for ln in SRC.read_text(encoding="utf-8").split("\n")]
    doc = Document()
    setup(doc)

    para(doc, lines[0].strip(), size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=11)
    para(doc, "First Author, Second Author", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "Department, Institution, City, Taiwan", size=9, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "author@institution.edu.tw", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    section, n_para, done = "", 0, set()
    for ln in lines[1:]:
        s = ln.strip()
        if not s:
            continue

        if s.startswith("ABSTRACT:"):
            p = doc.add_paragraph()
            a = p.add_run("Abstract. ")
            a.bold, a.font.size = True, Pt(9)
            b = p.add_run(s[len("ABSTRACT:"):].strip())
            b.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(8)
            continue
        if s.startswith("INDEX TERMS:"):
            p = doc.add_paragraph()
            a = p.add_run("Keywords: ")
            a.bold, a.font.size = True, Pt(9)
            b = p.add_run(s[len("INDEX TERMS:"):].strip())
            b.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(14)
            continue
        if s.startswith("## "):
            section = s[3:].strip()
            n_para = 0
            para(doc, section, size=12, bold=True, before=13, after=5)
            continue
        if s.startswith("### "):
            para(doc, s[4:].strip(), size=10, bold=True, italic=True, before=9, after=3)
            continue

        para(doc, s, indent=5 if n_para else 0, after=0)
        n_para += 1

        for prefix, at, kind, num in PLAN:
            if (kind, num) in done or not section.startswith(prefix) or n_para < at:
                continue
            done.add((kind, num))
            if kind == "fig":
                figure(doc, num)
            elif num == 1:
                table_one(doc)
            else:
                table_two(doc)

    missing = [(k, n) for _, _, k, n in PLAN if (k, n) not in done]
    for k, n in missing:
        print(f"  WARNING: {k} {n} was never placed", file=sys.stderr)

    target = OUT
    try:
        doc.save(target)
    except PermissionError:
        target = OUT.with_name(OUT.stem + "_new" + OUT.suffix)
        doc.save(target)
        print(f"NOTE: {OUT.name} is open in Word; wrote {target.name} instead.")
    print(f"wrote {target}  ({len(PLAN) - len(missing)}/{len(PLAN)} floats placed)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
